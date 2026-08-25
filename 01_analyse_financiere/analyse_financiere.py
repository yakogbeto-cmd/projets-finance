"""
Analyseur de sante financiere d'une entreprise
================================================
Calcule automatiquement les grands ratios financiers d'une entreprise a partir
de ses etats financiers, decompose le ROE facon DuPont, et exporte une fiche
d'analyse (Excel + graphique).

Deux modes :
  - Sans argument        -> utilise donnees_exemple.csv (marche toujours, pour demo)
  - --ticker AAPL / MC.PA -> recupere les vraies donnees via Yahoo Finance (yfinance)

Auteur : Yvan Akogbeto
Usage  : py analyse_financiere.py            (donnees d'exemple)
         py analyse_financiere.py --ticker MC.PA   (LVMH, donnees live)
"""

import argparse
import sys
import pandas as pd
import matplotlib
matplotlib.use("Agg")  # pas de fenetre, on sauvegarde en image
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage


# ---------------------------------------------------------------------------
# 1. Recuperation des donnees
# ---------------------------------------------------------------------------
def charger_depuis_csv(chemin="donnees_exemple.csv"):
    """Charge les postes comptables depuis le CSV d'exemple (en millions d'euros)."""
    df = pd.read_csv(chemin, index_col="poste")
    annee = df.columns[0]  # l'annee la plus recente
    d = df[annee].to_dict()
    d["_libelle"] = f"Societe Exemple SA (exercice {annee})"
    return d


def _cherche_ligne(df, candidats):
    """Cherche dans l'index d'un DataFrame yfinance la premiere ligne qui matche."""
    for nom in candidats:
        for idx in df.index:
            if nom.lower() in str(idx).lower():
                val = df.loc[idx].iloc[0]
                if pd.notna(val):
                    return float(val)
    return None


def charger_depuis_yahoo(ticker):
    """Recupere les postes comptables reels via yfinance (en devise de l'entreprise)."""
    import yfinance as yf
    t = yf.Ticker(ticker)
    inc, bs = t.income_stmt, t.balance_sheet
    if inc is None or inc.empty or bs is None or bs.empty:
        raise RuntimeError("Donnees indisponibles pour ce ticker.")
    d = {
        "chiffre_affaires": _cherche_ligne(inc, ["Total Revenue", "Operating Revenue"]),
        "resultat_net":     _cherche_ligne(inc, ["Net Income"]),
        "total_actif":      _cherche_ligne(bs, ["Total Assets"]),
        "capitaux_propres": _cherche_ligne(bs, ["Stockholders Equity", "Common Stock Equity", "Total Equity"]),
        "actif_courant":    _cherche_ligne(bs, ["Current Assets"]),
        "passif_courant":   _cherche_ligne(bs, ["Current Liabilities"]),
        "stocks":           _cherche_ligne(bs, ["Inventory"]) or 0.0,
        "dette_totale":     _cherche_ligne(bs, ["Total Debt"]) or 0.0,
    }
    manquants = [k for k, v in d.items() if v is None]
    if manquants:
        raise RuntimeError(f"Postes manquants pour {ticker}: {manquants}")
    d["_libelle"] = f"{t.info.get('longName', ticker)} ({ticker})"
    return d


# ---------------------------------------------------------------------------
# 2. Calcul des ratios
# ---------------------------------------------------------------------------
def calculer_ratios(d):
    """Retourne un dictionnaire {nom_ratio: (valeur, unite, interpretation)}."""
    ca, rn = d["chiffre_affaires"], d["resultat_net"]
    ta, cp = d["total_actif"], d["capitaux_propres"]
    ac, pc = d["actif_courant"], d["passif_courant"]
    stk, dette = d["stocks"], d["dette_totale"]

    r = {}
    # Liquidite : capacite a honorer les dettes court terme
    r["Ratio de liquidite generale"] = (ac / pc, "x", "actif courant / passif courant (>1 = sain)")
    r["Ratio de liquidite reduite"]  = ((ac - stk) / pc, "x", "(actif courant - stocks) / passif courant")
    # Endettement / structure
    r["Ratio d'endettement (D/E)"]   = (dette / cp, "x", "dette totale / capitaux propres")
    r["Autonomie financiere"]        = (cp / ta * 100, "%", "capitaux propres / total actif")
    # Rentabilite
    r["Marge nette"]                 = (rn / ca * 100, "%", "resultat net / chiffre d'affaires")
    r["ROA (rentabilite actif)"]     = (rn / ta * 100, "%", "resultat net / total actif")
    r["ROE (rentabilite fonds propres)"] = (rn / cp * 100, "%", "resultat net / capitaux propres")
    return r


def decomposer_dupont(d):
    """Decompose le ROE = marge nette x rotation de l'actif x levier financier."""
    marge = d["resultat_net"] / d["chiffre_affaires"]
    rotation = d["chiffre_affaires"] / d["total_actif"]
    levier = d["total_actif"] / d["capitaux_propres"]
    roe = marge * rotation * levier
    return {
        "Marge nette": marge * 100,
        "Rotation de l'actif": rotation,
        "Levier financier": levier,
        "= ROE reconstitue": roe * 100,
    }


# ---------------------------------------------------------------------------
# 3. Sorties : console, graphique, Excel
# ---------------------------------------------------------------------------
def afficher_console(libelle, ratios, dupont):
    print("=" * 60)
    print(f"  ANALYSE FINANCIERE — {libelle}")
    print("=" * 60)
    for nom, (val, unite, interp) in ratios.items():
        print(f"  {nom:<38} {val:>7.2f} {unite:<2}  ({interp})")
    print("-" * 60)
    print("  DECOMPOSITION DE DUPONT DU ROE")
    for nom, val in dupont.items():
        print(f"    {nom:<32} {val:>7.2f}")
    print("=" * 60)


def graphique_ratios(ratios, chemin="ratios.png"):
    noms = [n for n, (v, u, i) in ratios.items() if u == "%"]
    vals = [ratios[n][0] for n in noms]
    plt.figure(figsize=(7, 4))
    barres = plt.bar(noms, vals, color="#274060")
    plt.title("Ratios de rentabilite et de structure (%)")
    plt.ylabel("%")
    plt.xticks(rotation=25, ha="right", fontsize=8)
    for b, v in zip(barres, vals):
        plt.text(b.get_x() + b.get_width() / 2, v, f"{v:.1f}", ha="center", va="bottom", fontsize=8)
    plt.tight_layout()
    plt.savefig(chemin, dpi=130)
    plt.close()
    return chemin


def exporter_excel(libelle, ratios, dupont, img, chemin="fiche_analyse.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Analyse"
    titre = Font(bold=True, size=13, color="FFFFFF")
    entete = Font(bold=True, color="FFFFFF")
    fond = PatternFill("solid", fgColor="274060")

    ws["A1"] = f"Analyse financiere — {libelle}"
    ws["A1"].font = titre
    ws["A1"].fill = fond
    ws.merge_cells("A1:C1")

    ws.append([])
    ws.append(["Ratio", "Valeur", "Definition"])
    for c in ws[3]:
        c.font, c.fill = entete, fond
    for nom, (val, unite, interp) in ratios.items():
        ws.append([nom, f"{val:.2f} {unite}", interp])

    ws.append([])
    ligne = ws.max_row + 1
    ws.cell(ligne, 1, "Decomposition de Dupont du ROE").font = Font(bold=True)
    for nom, val in dupont.items():
        ws.append([nom, f"{val:.2f}"])

    for col, larg in {"A": 40, "B": 14, "C": 46}.items():
        ws.column_dimensions[col].width = larg

    ws.add_image(XLImage(img), f"A{ws.max_row + 2}")
    wb.save(chemin)
    return chemin


# ---------------------------------------------------------------------------
# 4. Programme principal
# ---------------------------------------------------------------------------
def main():
    p = argparse.ArgumentParser(description="Analyse financiere d'une entreprise.")
    p.add_argument("--ticker", help="Ticker Yahoo Finance (ex: MC.PA, AAPL). Sinon donnees d'exemple.")
    args = p.parse_args()

    if args.ticker:
        try:
            d = charger_depuis_yahoo(args.ticker)
        except Exception as e:
            print(f"[!] Echec donnees live ({e}). Bascule sur les donnees d'exemple.", file=sys.stderr)
            d = charger_depuis_csv()
    else:
        d = charger_depuis_csv()

    ratios = calculer_ratios(d)
    dupont = decomposer_dupont(d)
    afficher_console(d["_libelle"], ratios, dupont)
    img = graphique_ratios(ratios)
    xlsx = exporter_excel(d["_libelle"], ratios, dupont, img)
    print(f"\n[OK] Graphique : {img}")
    print(f"[OK] Fiche Excel : {xlsx}")


if __name__ == "__main__":
    main()
