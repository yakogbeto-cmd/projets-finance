"""
Tableau de bord de tresorerie + analyse d'ecarts budget/realise
================================================================
Genere un classeur Excel professionnel (formules vivantes, pas des valeurs figees)
avec trois feuilles :
  1. Tresorerie   : prevision mensuelle (solde initial, encaissements,
                    decaissements, flux net, solde final) sur 12 mois.
  2. Budget_Realise : analyse d'ecarts par poste, avec ecart en euros, en %,
                    et un statut Favorable/Defavorable calcule automatiquement
                    (une charge en dessous du budget = favorable ; un produit
                    au-dessus du budget = favorable).
  3. Dashboard    : graphiques (evolution du solde, budget vs realise).

Comme les cellules contiennent de VRAIES formules Excel, on peut changer une
hypothese (un encaissement, un budget) et tout le tableau se recalcule.

Auteur : Yvan Akogbeto
Usage  : py generer_tableau_tresorerie.py
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference

BLEU = "274060"
BLANC = "FFFFFF"
GRIS = "EEF1F5"

titre_font = Font(bold=True, size=13, color=BLANC)
entete_font = Font(bold=True, color=BLANC)
fond_bleu = PatternFill("solid", fgColor=BLEU)
fond_gris = PatternFill("solid", fgColor=GRIS)
bord = Border(*[Side(style="thin", color="CCCCCC")] * 4)
euro = '#,##0 "€"'
pourcent = '0.0 "%"'


def style_entete(ws, ligne, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(ligne, c)
        cell.font, cell.fill = entete_font, fond_bleu
        cell.alignment = Alignment(horizontal="center")


# ---------------------------------------------------------------------------
# Feuille 1 : Tresorerie previsionnelle
# ---------------------------------------------------------------------------
def feuille_tresorerie(wb):
    ws = wb.active
    ws.title = "Tresorerie"
    mois = ["Janvier", "Fevrier", "Mars", "Avril", "Mai", "Juin",
            "Juillet", "Aout", "Septembre", "Octobre", "Novembre", "Decembre"]
    # Donnees d'exemple (en euros) : encaissements / decaissements par mois
    encaissements = [120000, 118000, 135000, 128000, 140000, 150000,
                     110000, 95000, 138000, 145000, 152000, 160000]
    decaissements = [110000, 115000, 120000, 130000, 125000, 118000,
                     122000, 128000, 119000, 124000, 130000, 121000]
    solde_depart = 50000

    ws["A1"] = "Prevision de tresorerie — 12 mois (exemple)"
    ws["A1"].font, ws["A1"].fill = titre_font, fond_bleu
    ws.merge_cells("A1:F1")

    entetes = ["Mois", "Solde initial", "Encaissements", "Decaissements", "Flux net", "Solde final"]
    for c, h in enumerate(entetes, 1):
        ws.cell(3, c, h)
    style_entete(ws, 3, 6)

    for i, m in enumerate(mois):
        r = 4 + i
        ws.cell(r, 1, m)
        # Solde initial : depart le 1er mois, sinon solde final du mois precedent
        ws.cell(r, 2).value = solde_depart if i == 0 else f"=F{r-1}"
        ws.cell(r, 3, encaissements[i])
        ws.cell(r, 4, decaissements[i])
        ws.cell(r, 5).value = f"=C{r}-D{r}"          # flux net = encaiss - decaiss
        ws.cell(r, 6).value = f"=B{r}+E{r}"          # solde final = solde initial + flux net
        for c in range(2, 7):
            ws.cell(r, c).number_format = euro
        if i % 2:
            for c in range(1, 7):
                ws.cell(r, c).fill = fond_gris

    for col, larg in {"A": 12, "B": 14, "C": 15, "D": 15, "E": 13, "F": 14}.items():
        ws.column_dimensions[col].width = larg
    return ws, len(mois)


# ---------------------------------------------------------------------------
# Feuille 2 : Budget vs Realise (analyse d'ecarts)
# ---------------------------------------------------------------------------
def feuille_budget(wb):
    ws = wb.create_sheet("Budget_Realise")
    # (poste, type, budget, realise)  type = "Produit" ou "Charge"
    lignes = [
        ("Chiffre d'affaires", "Produit", 1500000, 1575000),
        ("Autres produits",    "Produit", 60000,   52000),
        ("Achats",             "Charge",  620000,  655000),
        ("Masse salariale",    "Charge",  480000,  472000),
        ("Charges externes",   "Charge",  180000,  176000),
        ("Marketing",          "Charge",  90000,   105000),
        ("Amortissements",     "Charge",  70000,   70000),
    ]

    ws["A1"] = "Analyse d'ecarts — Budget vs Realise (exemple)"
    ws["A1"].font, ws["A1"].fill = titre_font, fond_bleu
    ws.merge_cells("A1:F1")

    entetes = ["Poste", "Type", "Budget", "Realise", "Ecart (€)", "Ecart (%)", "Statut"]
    for c, h in enumerate(entetes, 1):
        ws.cell(3, c, h)
    style_entete(ws, 3, 7)

    for i, (poste, typ, bud, rea) in enumerate(lignes):
        r = 4 + i
        ws.cell(r, 1, poste)
        ws.cell(r, 2, typ)
        ws.cell(r, 3, bud).number_format = euro
        ws.cell(r, 4, rea).number_format = euro
        ws.cell(r, 5).value = f"=D{r}-C{r}"                 # ecart en euros
        ws.cell(r, 5).number_format = euro
        ws.cell(r, 6).value = f"=IF(C{r}=0,0,(D{r}-C{r})/C{r}*100)"  # ecart en %
        ws.cell(r, 6).number_format = pourcent
        # Statut : Produit favorable si realise > budget ; Charge favorable si realise < budget
        ws.cell(r, 7).value = (
            f'=IF(B{r}="Produit",IF(D{r}>=C{r},"Favorable","Defavorable"),'
            f'IF(D{r}<=C{r},"Favorable","Defavorable"))'
        )
        if i % 2:
            for c in range(1, 8):
                ws.cell(r, c).fill = fond_gris

    for col, larg in {"A": 20, "B": 10, "C": 13, "D": 13, "E": 12, "F": 11, "G": 13}.items():
        ws.column_dimensions[col].width = larg
    return ws, len(lignes)


# ---------------------------------------------------------------------------
# Feuille 3 : Dashboard (graphiques)
# ---------------------------------------------------------------------------
def feuille_dashboard(wb, n_mois, n_postes):
    ws = wb.create_sheet("Dashboard")
    ws["A1"] = "Tableau de bord"
    ws["A1"].font, ws["A1"].fill = titre_font, fond_bleu
    ws.merge_cells("A1:H1")

    # Courbe du solde final de tresorerie
    tr = wb["Tresorerie"]
    courbe = LineChart()
    courbe.title = "Evolution du solde de tresorerie"
    courbe.y_axis.title = "Solde (€)"
    donnees = Reference(tr, min_col=6, min_row=3, max_row=3 + n_mois)  # colonne Solde final + entete
    cats = Reference(tr, min_col=1, min_row=4, max_row=3 + n_mois)     # mois
    courbe.add_data(donnees, titles_from_data=True)
    courbe.set_categories(cats)
    courbe.height, courbe.width = 8, 16
    ws.add_chart(courbe, "A3")

    # Barres Budget vs Realise
    bg = wb["Budget_Realise"]
    barres = BarChart()
    barres.title = "Budget vs Realise par poste"
    barres.y_axis.title = "€"
    d = Reference(bg, min_col=3, max_col=4, min_row=3, max_row=3 + n_postes)  # Budget + Realise
    c = Reference(bg, min_col=1, min_row=4, max_row=3 + n_postes)             # postes
    barres.add_data(d, titles_from_data=True)
    barres.set_categories(c)
    barres.height, barres.width = 9, 18
    ws.add_chart(barres, "A20")
    return ws


def main():
    wb = Workbook()
    _, n_mois = feuille_tresorerie(wb)
    _, n_postes = feuille_budget(wb)
    feuille_dashboard(wb, n_mois, n_postes)
    sortie = "tableau_bord_tresorerie.xlsx"
    wb.save(sortie)
    print(f"[OK] Classeur genere : {sortie}")
    print("     Feuilles : Tresorerie | Budget_Realise | Dashboard")
    print("     Les cellules contiennent de vraies formules Excel (recalcul automatique).")


if __name__ == "__main__":
    main()
